import os
import uuid
import json
import hashlib
from pathlib import Path
from datetime import datetime, timezone
from typing import List, Tuple, Optional, Dict, Any
from sqlalchemy.orm import Session
from sqlalchemy import select, desc

from core.config import settings
from models import Document, DocumentChunk, BackgroundJob
from services.embedding_service import get_embedding

def calculate_checksum(file_bytes: bytes) -> str:
    """Computes SHA-256 checksum of file bytes for deduplication."""
    return hashlib.sha256(file_bytes).hexdigest()

def extract_text_from_file(file_path: Path, filename: str) -> List[Tuple[str, Dict[str, Any]]]:
    """
    Extracts text and page/section metadata from supported formats.
    Returns list of (text_content, metadata_dict).
    """
    ext = filename.rsplit(".", 1)[-1].lower() if "." in filename else ""
    results = []

    try:
        if ext == "pdf":
            import pypdf
            reader = pypdf.PdfReader(str(file_path))
            for page_idx, page in enumerate(reader.pages):
                text = page.extract_text() or ""
                if text.strip():
                    results.append((text.strip(), {"page": page_idx + 1, "source": filename}))

        elif ext == "docx":
            import docx
            doc = docx.Document(str(file_path))
            full_text = []
            for p in doc.paragraphs:
                if p.text.strip():
                    full_text.append(p.text.strip())
            for table in doc.tables:
                for row in table.rows:
                    row_text = " | ".join(cell.text.strip() for cell in row.cells if cell.text.strip())
                    if row_text:
                        full_text.append(row_text)
            combined = "\n".join(full_text)
            if combined:
                results.append((combined, {"source": filename}))

        elif ext in ["xlsx", "csv"]:
            if ext == "xlsx":
                import openpyxl
                wb = openpyxl.load_workbook(str(file_path), data_only=True)
                for sheet in wb.sheetnames:
                    ws = wb[sheet]
                    sheet_rows = []
                    for row in ws.iter_rows(values_only=True):
                        row_vals = [str(v) for v in row if v is not None]
                        if row_vals:
                            sheet_rows.append(" | ".join(row_vals))
                    if sheet_rows:
                        results.append(("\n".join(sheet_rows), {"sheet": sheet, "source": filename}))
            else:
                with open(file_path, "r", encoding="utf-8", errors="ignore") as f:
                    content = f.read()
                    if content.strip():
                        results.append((content.strip(), {"source": filename}))

        else:
            # txt, md, json, or other text formats
            with open(file_path, "r", encoding="utf-8", errors="ignore") as f:
                content = f.read()
                if content.strip():
                    results.append((content.strip(), {"source": filename}))

    except Exception as e:
        raise ValueError(f"Failed to extract text from {filename}: {e}")

    return results

def chunk_text(text: str, chunk_size: int = 900, overlap: int = 150) -> List[str]:
    """Splits long text into overlapping chunks respecting sentence/paragraph boundaries."""
    if len(text) <= chunk_size:
        return [text]

    chunks = []
    start = 0
    while start < len(text):
        end = start + chunk_size
        if end >= len(text):
            chunks.append(text[start:].strip())
            break

        # Look for sentence boundary near end
        cut = text.rfind(". ", start, end)
        if cut == -1 or cut <= start + (chunk_size // 2):
            cut = text.rfind("\n", start, end)
        if cut == -1 or cut <= start + (chunk_size // 2):
            cut = text.rfind(" ", start, end)
        if cut == -1 or cut <= start:
            cut = end

        chunk = text[start:cut].strip()
        if chunk:
            chunks.append(chunk)
        start = cut - overlap if (cut - overlap) > start else cut

    return chunks

def ingest_document(
    db: Session,
    user_id: uuid.UUID,
    filename: str,
    content_type: str,
    file_bytes: bytes
) -> Document:
    """Full document ingestion pipeline: validate -> save -> extract -> chunk -> embed -> index."""
    # 1. Validation
    ext = filename.rsplit(".", 1)[-1].lower() if "." in filename else ""
    if ext not in settings.ALLOWED_EXTENSIONS:
        raise ValueError(f"Unsupported file type .{ext}. Allowed: {', '.join(settings.ALLOWED_EXTENSIONS)}")

    max_bytes = settings.MAX_UPLOAD_SIZE_MB * 1024 * 1024
    if len(file_bytes) > max_bytes:
        raise ValueError(f"File size ({len(file_bytes)/1024/1024:.1f}MB) exceeds {settings.MAX_UPLOAD_SIZE_MB}MB limit.")

    checksum = calculate_checksum(file_bytes)

    # Check for existing duplicate for this user
    existing = db.scalar(
        select(Document).where(Document.user_id == user_id, Document.checksum == checksum)
    )
    if existing and existing.processing_status == "indexed":
        return existing

    # 2. Store on disk
    user_upload_dir = settings.UPLOADS_DIR / str(user_id)
    user_upload_dir.mkdir(parents=True, exist_ok=True)
    doc_id = uuid.uuid4()
    safe_filename = f"{doc_id}_{Path(filename).name}"
    storage_path = user_upload_dir / safe_filename

    with open(storage_path, "wb") as f:
        f.write(file_bytes)

    # 3. Create document record
    now = datetime.now(timezone.utc)
    doc = Document(
        id=doc_id,
        user_id=user_id,
        filename=filename,
        content_type=content_type,
        file_size=len(file_bytes),
        checksum=checksum,
        storage_path=str(storage_path),
        processing_status="processing",
        created_at=now
    )
    db.add(doc)
    db.commit()

    # 4. Extract and chunk
    try:
        sections = extract_text_from_file(storage_path, filename)
        if not sections:
            doc.processing_status = "failed"
            doc.error_message = "No readable text could be extracted from this document."
            db.commit()
            return doc

        chunk_counter = 0
        for section_text, sec_meta in sections:
            chunks = chunk_text(section_text)
            for c in chunks:
                if not c.strip():
                    continue
                # 5. Local 768-dim Embedding
                emb = get_embedding(c)
                chunk_obj = DocumentChunk(
                    id=uuid.uuid4(),
                    document_id=doc.id,
                    user_id=user_id,
                    chunk_index=chunk_counter,
                    text_content=c,
                    metadata_json=json.dumps(sec_meta),
                    embedding_json=json.dumps(emb),
                    created_at=now
                )
                db.add(chunk_obj)
                chunk_counter += 1

        doc.processing_status = "indexed"
        db.commit()
        db.refresh(doc)
        return doc

    except Exception as e:
        doc.processing_status = "failed"
        doc.error_message = str(e)
        db.commit()
        raise

def delete_document(db: Session, user_id: uuid.UUID, document_id: uuid.UUID) -> bool:
    """Deletes document record, cascading chunks, and storage file."""
    doc = db.scalar(
        select(Document).where(Document.id == document_id, Document.user_id == user_id)
    )
    if not doc:
        return False

    # Remove file on disk if exists
    try:
        if doc.storage_path and os.path.exists(doc.storage_path):
            os.remove(doc.storage_path)
    except Exception:
        pass

    db.delete(doc)
    db.commit()
    return True
